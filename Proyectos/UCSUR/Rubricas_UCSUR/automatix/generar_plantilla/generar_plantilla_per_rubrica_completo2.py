import os
import re
import pandas as pd
from docx import Document
from openpyxl import load_workbook

# === CONFIGURACIÓN ===
base_folder_path = r"C:\Users\manue\Downloads\LAST"
error_log = []
validation_log = []
document_log = []

# === MAPA DE EQUIVALENCIAS (ampliado) ===
mapa_equivalencias = {
    "EC1": "Evaluación Continua 1", "EC2": "Evaluación Continua 2", 
    "EC3": "Evaluación Continua 3", "EC4": "Evaluación Continua 4",
    "EF": "Evaluación Final", "EP": "Evaluación Parcial", 
    "ED": "Evaluación Diagnóstica",
    "CONTINUA1": "Evaluación Continua 1", "CONTINUA2": "Evaluación Continua 2", 
    "CONTINUA3": "Evaluación Continua 3", "CONTINUA4": "Evaluación Continua 4", 
    "DIAGNOSTICA": "Evaluación Diagnóstica", 
    "FINAL": "Evaluación Final", "PARCIAL": "Evaluación Parcial",
    "EVFINAL": "Evaluación Final (ES-FINAL)", "EV DIAG": "Evaluación Diagnóstica (Ficha)"
}

def detectar_tipo_evaluacion(nombre_archivo):
    nombre_upper = nombre_archivo.upper()
    for clave in mapa_equivalencias.keys():
        if clave in nombre_upper:
            return mapa_equivalencias[clave]
    return None

# === FUNCIONES DE PARSING MEJORADAS ===
def normalizar_puntaje(valor):
    if isinstance(valor, str):
        valor = valor.strip().replace(",", ".")
        # Manejo de casos especiales como en la versión original
        if valor.endswith("pts"):
            valor = valor[:-3].strip()
        elif valor.endswith("puntos"):
            valor = valor[:-6].strip()
    try:
        valor_float = float(valor)
        # Conversiones especiales como en la versión original
        if valor_float == 75:
            return 3.75
        elif valor_float == 50:
            return 5.0
        elif valor_float == 30:
            return 3.0
        return valor_float
    except (ValueError, TypeError):
        return None

def extraer_puntaje_desde_texto(texto):
    """
    Función para extraer solo el puntaje desde un texto
    (Usada en el procesamiento de tablas con formato de dos filas)
    """
    texto = texto.strip()
    # Busca el primer número en el texto
    match = re.search(r"(\d+[.,]?\d*)", texto)
    if match:
        return normalizar_puntaje(match.group(1))
    return None

def extraer_puntaje_y_descripcion(texto):
    """
    Versión mejorada que maneja diferentes formatos de descripción:
    1. Si el texto es solo el puntaje (ej: "6 pts"), devuelve (puntaje, "")
    2. Si contiene descripción y puntaje, extrae ambos
    3. Si no puede extraer puntaje, devuelve (None, texto_completo)
    """
    texto = texto.strip()
    
    # Caso 1: Solo el puntaje (ej: "6 pts" o "5puntos")
    solo_puntaje = re.fullmatch(r"^\d+[.,]?\d*\s*(?:pts|puntos)?$", texto, re.IGNORECASE)
    if solo_puntaje:
        puntaje = normalizar_puntaje(texto)
        return puntaje, ""
    
    # Caso 2: Descripción primero, luego puntaje (ej: "Excelente dominio 6 pts")
    match_final = re.search(r"(.*?)\s*(\d+[.,]?\d*)\s*(?:pts|puntos)?$", texto, re.IGNORECASE)
    if match_final:
        descripcion = match_final.group(1).strip()
        puntaje = normalizar_puntaje(match_final.group(2))
        return puntaje, descripcion
    
    # Caso 3: Puntaje primero, luego descripción (ej: "6 pts - Buen dominio")
    match_inicio = re.search(r"^(\d+[.,]?\d*)\s*(?:pts|puntos)?\s*[-–]?\s*(.*)", texto, re.IGNORECASE)
    if match_inicio:
        puntaje = normalizar_puntaje(match_inicio.group(1))
        descripcion = match_inicio.group(2).strip()
        return puntaje, descripcion
    
    # Caso 4: Buscar puntaje incrustado (ej: "Dominio (6 pts) del tema")
    match_incrustado = re.search(r"(.*?)\s*\(?\s*(\d+[.,]?\d*)\s*(?:pts|puntos)\s*\)?(.*)", texto, re.IGNORECASE)
    if match_incrustado:
        parte1 = match_incrustado.group(1).strip()
        parte2 = match_incrustado.group(3).strip()
        descripcion = f"{parte1} {parte2}".strip()
        puntaje = normalizar_puntaje(match_incrustado.group(2))
        return puntaje, descripcion
    
    # Caso 5: No se encontró patrón de puntaje, devolver todo como descripción
    return None, texto

def es_fila_cabecera(texto):
    return bool(re.search(r"ASPECTO|CRITERIO|EVALUAR|COMPETENCIAS|LOGRO DESTACADO|LOGRADO|NO LOGRADO", texto.upper()))

def es_fila_resumen(texto):
    return bool(re.search(r"\d+\s*[-–]\s*\d+\s*(?:pts|puntos)|TOTAL|PUNTAJE", texto.upper()))

# === PROCESAMIENTO DE DOCUMENTO WORD MEJORADO ===
def procesar_documento(word_file_path):
    try:
        doc = Document(word_file_path)
        evaluacion_data = []
        
        for table in doc.tables:
            # Intentar detectar si es una tabla de dos filas por criterio
            es_formato_dos_filas = False
            if len(table.rows) >= 2:
                primera_fila = table.rows[0].cells[0].text.strip().upper()
                segunda_fila = table.rows[1].cells[0].text.strip().upper()
                if "ASPECTO" not in primera_fila and "CRITERIO" not in primera_fila:
                    es_formato_dos_filas = True
            
            if es_formato_dos_filas:
                # Procesamiento para tablas con descripción y puntaje en filas separadas
                for i in range(0, len(table.rows) - 1, 2):
                    desc_row = table.rows[i]
                    puntaje_row = table.rows[i + 1]
                    
                    if len(desc_row.cells) < 4 or len(puntaje_row.cells) < 4:
                        continue
                        
                    criterio = desc_row.cells[0].text.strip()
                    desc_ld = desc_row.cells[1].text.strip()
                    desc_l = desc_row.cells[2].text.strip()
                    desc_nl = desc_row.cells[3].text.strip()
                    
                    p_ld = extraer_puntaje_desde_texto(puntaje_row.cells[1].text.strip())
                    p_l = extraer_puntaje_desde_texto(puntaje_row.cells[2].text.strip())
                    p_nl = extraer_puntaje_desde_texto(puntaje_row.cells[3].text.strip())
                    
                    if p_ld is not None and p_l is not None and p_nl is not None:
                        evaluacion_data.append([
                            criterio,
                            p_ld, "Logro Destacado", desc_ld,
                            p_l, "Logrado", desc_l,
                            p_nl, "No Logrado", desc_nl
                        ])
            else:
                # Procesamiento para tablas con un criterio por fila
                for row in table.rows:
                    if len(row.cells) < 4:
                        continue
                        
                    celda_0 = row.cells[0].text.strip()
                    celda_1 = row.cells[1].text.strip()
                    celda_2 = row.cells[2].text.strip()
                    celda_3 = row.cells[3].text.strip()
                    
                    if es_fila_cabecera(celda_0) or es_fila_resumen(celda_0):
                        continue
                        
                    if celda_0 and not any(x in celda_0.upper() for x in ["PTS", "PUNTOS", "TOTAL"]):
                        p_ld, desc_ld = extraer_puntaje_y_descripcion(celda_1)
                        p_l, desc_l = extraer_puntaje_y_descripcion(celda_2)
                        p_nl, desc_nl = extraer_puntaje_y_descripcion(celda_3)
                        
                        # Si no se extrajo descripción, usar el texto completo
                        desc_ld = desc_ld if desc_ld else celda_1
                        desc_l = desc_l if desc_l else celda_2
                        desc_nl = desc_nl if desc_nl else celda_3
                        
                        if p_ld is not None and p_l is not None and p_nl is not None:
                            evaluacion_data.append([
                                celda_0,
                                p_ld, "Logro Destacado", desc_ld,
                                p_l, "Logrado", desc_l,
                                p_nl, "No Logrado", desc_nl
                            ])

        if not evaluacion_data:
            error_log.append([word_file_path, "No se encontraron datos válidos en el documento"])
            return None
            
        df = pd.DataFrame(evaluacion_data, columns=[
            "ASPECTO / CRITERIO A EVALUAR", 
            "Puntaje de la calificación", "Título de la calificación", "Descripción de la calificación .1", 
            "Puntaje de la calificación.1", "Título de la calificación .1", "Descripción de la calificación .1.1", 
            "Puntaje de la calificación.2", "Título de la calificación .2", "Descripción de la calificación .2"
        ])
        
        # Validación de suma total con margen de error pequeño
        suma_ld = df["Puntaje de la calificación"].sum()
        if not (19.9 <= suma_ld <= 20.1):
            error_log.append([word_file_path, f"La suma de 'Logro Destacado' no es 20: {suma_ld:.2f}"])
            return None
            
        return df

    except Exception as e:
        error_log.append([word_file_path, f"Error al procesar el documento: {str(e)}"])
        return None

# === PROCESAMIENTO PRINCIPAL ===
def procesar_carpetas():
    for root, dirs, files in os.walk(base_folder_path):
        word_files = [f for f in files if f.endswith(".docx") and not f.startswith("~$")]
        if not word_files:
            continue
            
        folder_name = os.path.basename(root)
        output_excel_path = os.path.join(root, f"{folder_name}.xlsx")
        dfs_por_hoja = {}
        
        for filename in word_files:
            word_file_path = os.path.join(root, filename)
            tipo_eval = detectar_tipo_evaluacion(filename)
            
            if not tipo_eval:
                error_log.append([word_file_path, "Tipo de evaluación no reconocido"])
                document_log.append([filename, "Ignorado (tipo no reconocido)"])
                continue
                
            df = procesar_documento(word_file_path)
            if df is not None:
                sheet_name = tipo_eval[:31]
                
                if sheet_name in dfs_por_hoja:
                    counter = 1
                    while f"{sheet_name}_{counter}" in dfs_por_hoja:
                        counter += 1
                    sheet_name = f"{sheet_name}_{counter}"
                
                dfs_por_hoja[sheet_name] = df
                document_log.append([filename, f"Procesado correctamente -> {sheet_name}"])
            else:
                document_log.append([filename, "Procesado con errores"])
        
        if dfs_por_hoja:
            try:
                with pd.ExcelWriter(output_excel_path, engine='openpyxl') as writer:
                    for sheet_name, df in dfs_por_hoja.items():
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                print(f"✅ Archivo Excel generado: {output_excel_path}")
                
                validation_log.append([
                    output_excel_path,
                    ", ".join(word_files),
                    ", ".join(dfs_por_hoja.keys()),
                    "OK" if len(dfs_por_hoja) == len([f for f in document_log if "correctamente" in f[1]]) else "Algunos archivos no se procesaron"
                ])
            except Exception as e:
                error_log.append([output_excel_path, f"Error al generar Excel: {str(e)}"])

# === EJECUCIÓN Y REPORTES ===
if __name__ == "__main__":
    print("Iniciando procesamiento de documentos...")
    procesar_carpetas()
    
    # Generar reportes
    if error_log:
        pd.DataFrame(error_log, columns=["Archivo", "Error"]).to_excel(
            os.path.join(base_folder_path, "Errores_de_procesamiento.xlsx"), index=False)
        
    if validation_log:
        pd.DataFrame(validation_log, columns=["Archivo Excel", "Archivos Word", "Hojas en Excel", "Observación"]).to_excel(
            os.path.join(base_folder_path, "Resumen_de_validacion.xlsx"), index=False)
        
    if document_log:
        pd.DataFrame(document_log, columns=["Archivo", "Estado"]).to_excel(
            os.path.join(base_folder_path, "Resumen_de_procesamiento.xlsx"), index=False)
    
    print("\nResumen:")
    print(f"- Documentos procesados correctamente: {len([x for x in document_log if 'correctamente' in x[1]])}")
    print(f"- Documentos con errores: {len([x for x in document_log if 'errores' in x[1]])}")
    print(f"- Errores registrados: {len(error_log)}")
    print("\nProceso completado. Revise los archivos de reporte si hay errores.")