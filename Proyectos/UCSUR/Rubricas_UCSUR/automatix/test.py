import os
import re
import pandas as pd
from docx import Document
from openpyxl import load_workbook

# === CONFIGURACIÓN ===
base_folder_path = r"C:\Users\manue\OneDrive - Grupo Educad\3. Documentos validados 20252\Ingeniería Agroforestal y Ambiental\Manejo de Sistemas Agroforestales I"
error_log = []
validation_log = []
document_log = []

# === MAPA DE EQUIVALENCIAS ===
mapa_equivalencias = {
    "EC1": "Evaluación Continua 1", "EC2": "Evaluación Continua 2", "EC3": "Evaluación Continua 3", "EC4": "Evaluación Continua 4",
    "EF": "Evaluación Final", "EP": "Evaluación Parcial", "ED": "Evaluación Diagnóstica",
    "CONTINUA1": "Evaluación Continua 1", "CONTINUA2": "Evaluación Continua 2", "CONTINUA3": "Evaluación Continua 3",
    "CONTINUA4": "Evaluación Continua 4", "CONTINUAI": "Evaluación Continua I", "CONTINUAII": "Evaluación Continua II",
    "DIAGNOSTICA": "Evaluación Diagnóstica", "FINAL": "Evaluación Final", "PARCIAL": "Evaluación Parcial",
    "EVFINAL": "Evaluación Final (ES-FINAL)", "EV DIAG": "Evaluación Diagnóstica (Ficha)"
}

def detectar_tipo_evaluacion(nombre_archivo):
    for clave in mapa_equivalencias.keys():
        if clave in nombre_archivo.upper():
            return mapa_equivalencias[clave]
    return None

def normalizar_puntaje(valor):
    if isinstance(valor, str):
        valor = valor.strip().replace(",", ".")
    try:
        valor_float = float(valor)
        if valor_float == 75:
            return 3.75
        elif valor_float == 50:
            return 5.0
        elif valor_float == 30:
            return 3.0
        return valor_float
    except:
        return None

def extraer_puntaje_desde_texto(texto):
    posibles = re.findall(r"(\d+[.,]?\d*)", texto)
    valores_norm = [normalizar_puntaje(val) for val in posibles]
    validos = [v for v in valores_norm if v is not None]
    if validos:
        return validos[0]
    return 0

# ✅ Versión mejorada de detección de filas guía
def es_fila_resumen(*textos):
    patrones = [
        r"\d+\s*[-–]\s*\d+\s*pto?s?\.?",  # Ej: 16–20 ptos.
        r"\d+\s*[-–]\s*\d+\s*$",          # Ej: 0–12
    ]
    contador = sum(
        any(re.search(pat, texto.lower()) for pat in patrones)
        for texto in textos
    )
    return contador >= 2


def procesar_documento(word_file_path):
    try:
        doc = Document(word_file_path)
        tables = doc.tables
        evaluacion_data = []

        if not tables:
            error_log.append([word_file_path, "El documento no contiene tablas procesables"])
            return None

        for table in tables:
            rows = table.rows
            if len(rows) < 2:
                continue

            for i in range(1, len(rows), 2):
                if i + 1 >= len(rows):
                    continue

                desc_row = rows[i]
                puntaje_row = rows[i + 1]

                if len(desc_row.cells) < 4 or len(puntaje_row.cells) < 4:
                    continue

                criterio = desc_row.cells[0].text.strip()
                desc_ld = desc_row.cells[1].text.strip()
                desc_l = desc_row.cells[2].text.strip()
                desc_nl = desc_row.cells[3].text.strip()

                if criterio.upper() in ["ASPECTO / CRITERIO A EVALUAR", "ASPECTO O CRITERIO A EVALUAR"]:
                    continue

                if all(p.lower().startswith("total") for p in [criterio, desc_ld, desc_l, desc_nl]):
                    continue

                if es_fila_resumen(criterio, desc_ld, desc_l, desc_nl):
                    continue

                if not criterio or not desc_ld or not desc_l or not desc_nl:
                    continue

                p_ld = extraer_puntaje_desde_texto(puntaje_row.cells[1].text.strip()) or extraer_puntaje_desde_texto(desc_ld)
                p_l = extraer_puntaje_desde_texto(puntaje_row.cells[2].text.strip()) or extraer_puntaje_desde_texto(desc_l)
                p_nl = extraer_puntaje_desde_texto(puntaje_row.cells[3].text.strip()) or extraer_puntaje_desde_texto(desc_nl)

                evaluacion_data.append([
                    criterio,
                    p_ld, "Logro Destacado", desc_ld,
                    p_l, "Logrado", desc_l,
                    p_nl, "No Logrado", desc_nl
                ])

        if not evaluacion_data:
            error_log.append([word_file_path, "El documento no generó datos válidos"])
            return None

        df = pd.DataFrame(evaluacion_data, columns=[
            "ASPECTO / CRITERIO A EVALUAR", "Puntaje de la calificación", "Título de la calificación",
            "Descripción de la calificación .1", "Puntaje de la calificación.1", "Título de la calificación .1",
            "Descripción de la calificación .1.1", "Puntaje de la calificación.2", "Título de la calificación .2",
            "Descripción de la calificación .2"
        ])

        # ✅ EXCLUIR FILA FINAL SI ES RESUMEN
        if not df.empty:
            ultima_fila = df.iloc[-1]
            if es_fila_resumen(
                str(ultima_fila["ASPECTO / CRITERIO A EVALUAR"]),
                str(ultima_fila["Descripción de la calificación .1"]),
                str(ultima_fila["Descripción de la calificación .1.1"]),
                str(ultima_fila["Descripción de la calificación .2"])
            ):
                df = df[:-1]  # Excluir última fila si es resumen

        # Validar suma
        suma_ld = df["Puntaje de la calificación"].sum()
        if round(suma_ld, 2) != 20.0:
            error_log.append([word_file_path, f"La suma de 'Logro Destacado' no es 20: {suma_ld}"])
            return None

        return df

    except Exception as e:
        error_log.append([word_file_path, f"Error al procesar: {str(e)}"])
        return None

# === PROCESAMIENTO DE CARPETAS ===
for root, dirs, files in os.walk(base_folder_path):
    word_files = [f for f in files if f.endswith(".docx")]
    if word_files:
        folder_name = os.path.basename(root)
        output_excel_path = os.path.join(root, f"{folder_name}.xlsx")
        generated_sheets = []
        dfs_por_hoja = {}

        for filename in word_files:
            word_file_path = os.path.join(root, filename)
            df_reporte = procesar_documento(word_file_path)
            tipo_eval = detectar_tipo_evaluacion(filename)

            if tipo_eval is None:
                error_log.append([word_file_path, "No se detectó tipo de evaluación válido"])
                document_log.append([filename, "Ignorado (sin tipo de evaluación reconocido)"])
                continue

            sheet_name = tipo_eval[:31]
            if sheet_name in generated_sheets:
                counter = 1
                while f"{sheet_name}_{counter}" in generated_sheets:
                    counter += 1
                sheet_name = f"{sheet_name}_{counter}"

            if df_reporte is not None:
                dfs_por_hoja[sheet_name] = df_reporte
                generated_sheets.append(sheet_name)

            document_log.append([filename, "Procesado" if df_reporte is not None else "Ignorado"])

        if generated_sheets:
            with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
                for nombre_hoja, df in dfs_por_hoja.items():
                    df.to_excel(writer, sheet_name=nombre_hoja, index=False)
            print(f"✅ Reporte generado en: {output_excel_path}")

# === REPORTES GLOBALES ===
if error_log:
    pd.DataFrame(error_log, columns=["Archivo", "Descripción del Problema"]).to_excel(
        os.path.join(base_folder_path, "Reporte_Errores.xlsx"), index=False)

if validation_log:
    pd.DataFrame(validation_log, columns=["Archivo Excel", "Archivos Word", "Hojas en Excel", "Observación"]).to_excel(
        os.path.join(base_folder_path, "Reporte_Validaciones.xlsx"), index=False)
